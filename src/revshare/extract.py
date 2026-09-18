from __future__ import annotations

import csv
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from revshare.fx import valid_period

MONEY_RE = re.compile(
    r"(?<![\w.])(?P<amount>"
    r"\d{1,3}(?:[ \u00a0]\d{3})+(?:[.,]\d{1,2})?"
    r"|\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?"
    r"|\d+(?:[.,]\d{1,2})?"
    r")(?![\d\w])"
)
CURRENCY_RE = re.compile(r"\b(USD|EUR|AMD|RUB|KZT|AED|TRY|UAH|GBP|PLN|CZK)\b", re.I)
TOTAL_HINTS = (
    "total",
    "оборот",
    "выручка",
    "gmv",
    "turnover",
    "revenue",
    "revshare",
    "ревшер",
    "к оплате",
    "due",
    "share",
)


@dataclass
class AmountCandidate:
    amount: float
    raw: str
    currency: str | None = None
    label: str = ""
    confidence: str = "low"


@dataclass
class ExtractedReport:
    source: str
    partner_hint: str
    candidates: list[AmountCandidate] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    error: str | None = None

    def to_dict(self) -> dict:
        payload = asdict(self)
        return payload


@dataclass
class ParsedPartnerReport:
    turnover: float | None
    currency: str | None
    reported_share: float | None
    location_ids: list[str]
    location_source: str
    notes: list[str] = field(default_factory=list)


def _parse_amount(raw: str) -> float | None:
    cleaned = raw.strip().replace("\u00a0", " ").replace(" ", "")
    if cleaned.count(",") == 1 and cleaned.count(".") == 0:
        cleaned = cleaned.replace(",", ".")
    elif cleaned.count(".") > 1 and "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif cleaned.count(",") > 1:
        cleaned = cleaned.replace(",", "")
    elif "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


TURNOVER_HINTS = (
    "turnover",
    "revenue",
    "gmv",
    "amount",
    "оборот",
    "выручка",
    "виручка",
    "сумма",
)
SHARE_HINTS = (
    "revshare",
    "revenue share",
    "reward",
    "commission",
    "комиссия",
    "комісія",
    "altegio",
    "к оплате",
    "до сплати",
)
CURRENCY_SYMBOLS = {
    "$": "USD",
    "€": "EUR",
    "₸": "KZT",
    "₴": "UAH",
    "₽": "RUB",
}


def _currency_from_text(text: str, default: str | None) -> str | None:
    code = CURRENCY_RE.search(text)
    if code:
        return code.group(1).upper()
    for symbol, currency in CURRENCY_SYMBOLS.items():
        if symbol in text:
            return currency
    return default


def _value_from_cell(cell: str) -> float | None:
    if not cell.strip() or "%" in cell:
        return None
    match = MONEY_RE.search(cell)
    return _parse_amount(match.group("amount")) if match else None


def _signed_value_from_cell(cell: str) -> float | None:
    value = _value_from_cell(cell)
    if value is None:
        return None
    if cell.strip().startswith(("-", "−")):
        return -abs(value)
    return value


def _total_line_values(rows: list[list[str]]) -> tuple[float | None, float | None]:
    """Read explicit total rows, keeping turnover and Altegio share apart."""
    turnover = None
    share = None
    for row in rows:
        line = " ".join(cell for cell in row if cell.strip())
        if not _is_total_line(line):
            continue
        match = None
        for match in MONEY_RE.finditer(line):
            pass
        if match is None:
            continue
        value = _parse_amount(match.group("amount"))
        if value is None:
            continue
        folded = line.casefold()
        if any(hint in folded for hint in SHARE_HINTS):
            share = value
        else:
            turnover = value
    return turnover, share


def _labeled_value(rows: list[list[str]], hints: tuple[str, ...]) -> float | None:
    matches: list[float] = []
    for row in rows:
        for column, cell in enumerate(row):
            folded = cell.casefold()
            if not any(hint in folded for hint in hints):
                continue
            if _is_location_header(cell):
                continue
            own_value = _value_from_cell(cell)
            if own_value is not None:
                matches.append(own_value)
            for adjacent in row[column + 1 : column + 3]:
                value = _value_from_cell(adjacent)
                if value is not None:
                    matches.append(value)
                    break
    return matches[-1] if matches else None


def _column_total(rows: list[list[str]], hints: tuple[str, ...]) -> float | None:
    for header_index, row in enumerate(rows[:30]):
        for column, cell in enumerate(row):
            folded = cell.casefold().strip()
            if not any(hint in folded for hint in hints):
                continue
            if _is_location_header(cell):
                continue
            values: list[float] = []
            explicit_total = None
            for data_row in rows[header_index + 1 :]:
                if column >= len(data_row):
                    continue
                value = _value_from_cell(data_row[column])
                if value is None:
                    continue
                row_label = " ".join(data_row[:column]).casefold()
                if any(label in row_label for label in ("total", "итого", "всего", "разом")):
                    explicit_total = value
                else:
                    values.append(value)
            if explicit_total is not None:
                return explicit_total
            if values:
                return round(sum(values), 2)
    return None


def _standard_partner_workbook(
    path: Path, period: str | None = None
) -> ParsedPartnerReport | None:
    """Read the recommended Integration Hub partner report.

    Partners fill the input columns; formula cells may have no cached values,
    so totals are calculated here from amount, FX and contractual rate.
    """
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return None
    workbook = load_workbook(path, data_only=True, read_only=True)
    wanted = period_sheet_titles(period)
    sheets = [
        sheet
        for sheet in workbook.worksheets
        if not wanted or " ".join(sheet.title.split()).casefold() in wanted
    ] or list(workbook.worksheets)
    for sheet in sheets:
        rows = _sheet_rows(sheet)
        header_index = None
        columns: dict[str, int] = {}
        for index, row in enumerate(rows[:20]):
            for column, cell in enumerate(row):
                folded = " ".join(cell.split()).casefold()
                if _is_location_header(cell):
                    columns["location"] = column
                if (
                    "received by partner" in folded
                    or "получено партнёром" in folded
                ) and "report currency" not in folded and "валюта отчёта" not in folded:
                    columns["amount"] = column
                if folded.startswith("payment currency") or folded.startswith(
                    "валюта платежа"
                ):
                    columns["currency"] = column
                if "monthly avg rate" in folded or "среднемес. курс" in folded:
                    columns["fx"] = column
                if (
                    ("altegio" in folded and "rate" in folded)
                    or ("ставка" in folded and "altegio" in folded)
                ):
                    columns["rate"] = column
            if {"location", "amount", "currency"} <= columns.keys():
                header_index = index
                break
        if header_index is None:
            continue

        location_ids: list[str] = []
        turnover = 0.0
        share = 0.0
        used_rows = 0
        skipped_fx = 0
        report_currency = "USD"
        header_text = " ".join(rows[header_index])
        report_match = re.search(
            r"\((USD|EUR|AMD|RUB|KZT|AED|TRY|UAH|GBP|PLN|CZK)"
            r"(?:,\s*(?:report currency|валюта отч[её]та))?\)",
            header_text,
            re.I,
        )
        if report_match:
            report_currency = report_match.group(1).upper()

        for row in rows[header_index + 1 :]:
            line = " ".join(row)
            if _is_total_line(line):
                break
            location_cell = (
                _plain_id_cell(row[columns["location"]])
                if columns["location"] < len(row)
                else ""
            )
            location_match = LOCATION_COLUMN_ID_RE.search(location_cell)
            if not location_match:
                continue
            location_ids.append(location_match.group(1))
            amount_cell = row[columns["amount"]] if columns["amount"] < len(row) else ""
            amount = _signed_value_from_cell(amount_cell)
            if amount is None:
                continue
            currency_cell = (
                row[columns["currency"]] if columns["currency"] < len(row) else ""
            )
            payment_currency = _currency_from_text(currency_cell, report_currency)
            fx = 1.0
            if payment_currency != report_currency:
                fx_column = columns.get("fx")
                fx_cell = (
                    row[fx_column]
                    if fx_column is not None and fx_column < len(row)
                    else ""
                )
                fx = _value_from_cell(fx_cell) or 0.0
                if not fx:
                    skipped_fx += 1
                    continue
            converted = amount * fx
            turnover += converted
            rate = None
            if "rate" in columns and columns["rate"] < len(row):
                rate = _value_from_cell(row[columns["rate"]])
            if rate is not None:
                share += converted * rate
            used_rows += 1

        if not location_ids:
            continue
        notes = ["recommended_template"]
        if skipped_fx:
            notes.append("rows_missing_fx")
        if not used_rows:
            notes.append("turnover_not_found")
        return ParsedPartnerReport(
            turnover=round(turnover, 2) if used_rows else None,
            currency=report_currency,
            reported_share=round(share, 2) if share else None,
            location_ids=_unique(location_ids),
            location_source="recommended_template",
            notes=notes,
        )
    return None


def parse_partner_report(
    path: Path,
    rate: float | None = None,
    default_currency: str | None = None,
    period: str | None = None,
) -> ParsedPartnerReport:
    standard = _standard_partner_workbook(path, period)
    if standard is not None:
        if standard.reported_share is None and standard.turnover is not None and rate is not None:
            standard.reported_share = round(standard.turnover * rate, 2)
            standard.notes.append("share_calculated")
        return standard
    rows = _rows_from_file(path, period)
    text = "\n".join(" ".join(row) for row in rows)
    location_ids, location_source = location_ids_from_file(path, period)
    notes: list[str] = []

    turnover, reported_share = _total_line_values(rows)
    if turnover is None:
        turnover = _labeled_value(rows, TURNOVER_HINTS) or _column_total(
            rows, TURNOVER_HINTS
        )
    if reported_share is None:
        reported_share = _labeled_value(rows, SHARE_HINTS) or _column_total(
            rows, SHARE_HINTS
        )
    if turnover is None:
        notes.append("turnover_not_found")
    if reported_share is None and turnover is not None and rate is not None:
        reported_share = round(turnover * rate, 2)
        notes.append("share_calculated")
    if not location_ids:
        notes.append("locations_not_found")
    currency = _currency_from_text(text, default_currency)
    return ParsedPartnerReport(
        turnover=turnover,
        currency=currency,
        reported_share=reported_share,
        location_ids=location_ids,
        location_source=location_source,
        notes=notes,
    )


def amounts_from_text(text: str, default_currency: str | None = None) -> list[AmountCandidate]:
    found: list[AmountCandidate] = []
    currency_global = None
    currency_match = CURRENCY_RE.search(text or "")
    if currency_match:
        currency_global = currency_match.group(1).upper()
    for match in MONEY_RE.finditer(text or ""):
        raw = match.group("amount")
        value = _parse_amount(raw)
        if value is None:
            continue
        start = max(0, match.start() - 40)
        label = " ".join((text or "")[start : match.start()].split())
        confidence = "medium" if any(h in label.casefold() for h in TOTAL_HINTS) else "low"
        if value < 1:
            continue
        found.append(
            AmountCandidate(
                amount=value,
                raw=raw,
                currency=(currency_global or default_currency),
                label=label,
                confidence=confidence,
            )
        )
    hinted = [c for c in found if c.confidence == "medium"]
    return hinted or found


def extract_plain(path: Path) -> ExtractedReport:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return ExtractedReport(
        source=str(path),
        partner_hint=path.stem,
        candidates=amounts_from_text(text),
    )


def extract_csv(path: Path) -> ExtractedReport:
    rows: list[str] = []
    with path.open(newline="", encoding="utf-8", errors="ignore") as handle:
        for row in csv.reader(handle):
            rows.append(" ".join(row))
    text = "\n".join(rows)
    return ExtractedReport(
        source=str(path),
        partner_hint=path.stem,
        candidates=amounts_from_text(text),
    )


def extract_xlsx(path: Path) -> ExtractedReport:
    workbook = load_workbook(path, data_only=True, read_only=True)
    notes: list[str] = []
    candidates: list[AmountCandidate] = []
    partner_hint = path.stem
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            line = " ".join(cells)
            if not line.strip():
                continue
            for candidate in amounts_from_text(line):
                candidate.label = f"{sheet.title}: {candidate.label or line[:80]}"
                candidates.append(candidate)
    hinted = [c for c in candidates if c.confidence == "medium"]
    picked = hinted or candidates
    if len(picked) > 8:
        notes.append(f"{len(picked)} numeric candidates; showing hinted totals first")
    return ExtractedReport(
        source=str(path),
        partner_hint=partner_hint,
        candidates=picked[:12],
        notes=notes,
    )


def extract_pdf(path: Path) -> ExtractedReport:
    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    notes = []
    if not text.strip():
        notes.append("PDF has no extractable text; request Excel or a screenshot from the partner")
    return ExtractedReport(
        source=str(path),
        partner_hint=path.stem,
        candidates=amounts_from_text(text),
        notes=notes,
        error="empty pdf text" if not text.strip() else None,
    )


EXTRACTORS = {
    ".txt": extract_plain,
    ".md": extract_plain,
    ".csv": extract_csv,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".xls": extract_xlsx,
    ".pdf": extract_pdf,
}


def extract_file(path: Path) -> ExtractedReport:
    suffix = path.suffix.lower()
    extractor = EXTRACTORS.get(suffix)
    if extractor is None:
        return ExtractedReport(
            source=str(path),
            partner_hint=path.stem,
            error=f"unsupported type {suffix}",
            notes=["Forward as-is for manual review; screenshots are read in chat, not by this CLI"],
        )
    try:
        return extractor(path)
    except Exception as exc:  # noqa: BLE001 — intake must not crash the month run
        return ExtractedReport(
            source=str(path),
            partner_hint=path.stem,
            error=str(exc),
        )


LOCATION_HEADER_NAMES = {
    "id",
    "ids",
    "location",
    "locations",
    "location id",
    "location_id",
    "salon",
    "salon id",
    "salon_id",
    "branch",
    "branch id",
    "id altegio",
    "altegio id",
    "филиал",
    "id филиала",
    "локация",
    "точка",
    "заклад",
}
LOCATION_HEADER_PHRASES = (
    "location id",
    "location_id",
    "salon id",
    "salon_id",
    "id altegio",
    "altegio id",
    "id филиала",
)
TOTAL_LABELS = ("total", "totals", "итого", "всего", "разом")
LOCATION_ID_RE = re.compile(r"(?<![\d.,])(\d{4,8})(?:-alt)?(?![\d.,])", re.I)
LOCATION_COLUMN_ID_RE = re.compile(r"(?<![\d.,])(\d{2,12})(?:-alt)?(?![\d.,])", re.I)
FLOAT_ID_RE = re.compile(r"^\s*(\d+)[.,]0+\s*$")


def _plain_id_cell(cell: str) -> str:
    """Spreadsheets often store Location IDs as floats, so 1295757.0 is one id."""
    match = FLOAT_ID_RE.match(cell or "")
    return match.group(1) if match else (cell or "")


def _is_location_header(cell: str) -> bool:
    folded = " ".join(cell.split()).casefold().strip(" :#")
    if not folded or len(folded) > 40:
        return False
    if folded in LOCATION_HEADER_NAMES:
        return True
    return any(phrase in folded for phrase in LOCATION_HEADER_PHRASES)


def _is_total_line(text: str) -> bool:
    folded = text.casefold()
    return any(label in folded for label in TOTAL_LABELS)


def _is_tabular(rows: list[list[str]]) -> bool:
    return any(len([cell for cell in row if cell.strip()]) > 1 for row in rows)


def period_sheet_titles(period: str | None) -> set[str]:
    """Titles that can name the reporting month, so one month is read at a time."""
    checked = valid_period(period or "")
    if not checked:
        return set()
    year, month = checked.split("-")
    titles = {
        checked,
        f"{month}.{year}",
        f"{month}/{year}",
        f"{int(month)}.{year}",
        f"{year}/{month}",
    }
    from revshare.i18n import TRANSLATIONS

    for language in TRANSLATIONS.values():
        name = language.get(f"month_{month}")
        if name:
            titles.add(f"{name} {year}")
            titles.add(f"{name[:3]} {year}")
    return {title.casefold() for title in titles}


def _sheet_rows(sheet) -> list[list[str]]:
    return [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]


def _rows_from_file(path: Path, period: str | None = None) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        workbook = load_workbook(path, data_only=True, read_only=True)
        wanted = period_sheet_titles(period)
        if wanted:
            for sheet in workbook.worksheets:
                if " ".join(sheet.title.split()).casefold() in wanted:
                    return _sheet_rows(sheet)
        rows: list[list[str]] = []
        for sheet in workbook.worksheets:
            rows.extend(_sheet_rows(sheet))
        return rows
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8", errors="ignore") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [list(row) for row in csv.reader(handle, dialect)]
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return [[line] for line in text.splitlines()]
    return [[line] for line in path.read_text(encoding="utf-8", errors="ignore").splitlines()]


def location_ids_from_file(
    path: Path, period: str | None = None
) -> tuple[list[str], str]:
    """Read the client registry out of a partner file.

    Returns the ids and how they were found, so the submission can tell the
    partner whether their own file was understood.
    """
    try:
        rows = _rows_from_file(path, period)
    except Exception:  # noqa: BLE001 — a broken upload must not lose the report
        return [], "unreadable"

    if _is_tabular(rows):
        for index, row in enumerate(rows[:10]):
            for column, cell in enumerate(row):
                if not _is_location_header(cell):
                    continue
                found: list[str] = []
                for data_row in rows[index + 1 :]:
                    if column >= len(data_row):
                        continue
                    if _is_total_line(" ".join(data_row)):
                        continue
                    match = LOCATION_COLUMN_ID_RE.search(
                        _plain_id_cell(data_row[column])
                    )
                    if match:
                        found.append(match.group(1))
                if found:
                    return _unique(found), "column"

    scanned: list[str] = []
    for row in rows:
        line = " ".join(row)
        if _is_total_line(line):
            continue
        match = LOCATION_ID_RE.search(line)
        if match:
            scanned.append(match.group(1))
    return (_unique(scanned), "scan") if scanned else ([], "not_found")


def _unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def ingest_dir(folder: Path) -> list[ExtractedReport]:
    files = sorted(
        p
        for p in folder.iterdir()
        if p.is_file() and p.name not in {".gitkeep", "README.md"} and not p.name.startswith(".")
    )
    return [extract_file(path) for path in files]
