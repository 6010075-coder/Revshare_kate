"""Tests for the partner-facing XLSX report template.

Layout (post-2026-10 redesign):

- Sheet 1 — partner-language form, title embeds the reporting month so
  the parser recognises it (e.g. "Report (EN) — September 2026").
- Sheet 2 — "Пример отчёта" filled-in example the partner can study.
  The parser explicitly skips sheets whose title contains
  "пример" / "sample" / "example".

Both sheets share the same ten-column schema: location ID, client
name/email, subscription period, payment currency, received by partner
in payment currency, monthly avg rate to report currency, settlement
currency with Altegio, received by partner in report currency, Altegio
rate, due to Altegio.
"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook

from revshare.report_template import (
    COLUMN_WIDTHS,
    DATA_ROW_COUNT,
    DISCLAIMER_BODY_ROW,
    DISCLAIMER_HEADER_ROW,
    EXPECTED_HEADERS_EN,
    EXPECTED_HEADERS_RU,
    FIRST_DATA_ROW,
    LAST_DATA_ROW,
    TOTAL_ROW,
    TRANSLATIONS,
    build_report_template_xlsx,
)
from revshare.store import Store


def _build_partner(tmp_path: Path, **overrides):
    """Return ``(store, partner)`` so ``build_report_template_xlsx`` reads
    from the same store the partner was added to.
    """
    store = Store(tmp_path / "template.db")
    payload = {
        "name": "Acme Bookings, Inc.",
        "legal_name": "",
        "legal_country": "",
        "altegio_legal_entity": "",
        "contact_name": "Daniel Carter, CFO",
        "email": "billing@acme.io",
        "telegram": "",
        "report_channel": "telegram",
        "token": "tok",
        "rate": 0.2,
        "currency": "USD",
    }
    payload.update(overrides)
    pid = store.add_partner(**payload)
    return store, store.partner(pid)


def _primary_sheet_title(workbook, period: str, language: str) -> str:
    year, month = period.split("-")
    translations = TRANSLATIONS[language]
    base = translations["sheet_name"].split(" — ", 1)[0]
    from revshare.i18n import TRANSLATIONS as I18N

    month_name = I18N[language].get(f"month_{int(month):02d}", "")
    return f"{base} — {month_name} {year}"


def test_build_report_template_returns_bytes(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09")
    assert isinstance(blob, bytes)
    assert len(blob) > 1000
    assert blob[:2] == b"PK"


def test_ru_partner_gets_ru_form_sheet_then_example(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    assert workbook.sheetnames[0] == _primary_sheet_title(workbook, "2026-09", "ru")
    assert workbook.sheetnames[1] == TRANSLATIONS["ru"]["example_sheet_name"]


def test_en_partner_gets_en_form_sheet_then_example(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    assert workbook.sheetnames[0] == _primary_sheet_title(workbook, "2026-09", "en")
    assert workbook.sheetnames[1] == TRANSLATIONS["ru"]["example_sheet_name"]


def test_ukrainian_partner_falls_back_to_ru_form(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="uk")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    assert workbook.sheetnames[0] == _primary_sheet_title(workbook, "2026-09", "ru")


def test_polish_partner_falls_back_to_ru_form(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="pl")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    assert workbook.sheetnames[0] == _primary_sheet_title(workbook, "2026-09", "ru")


def test_ru_sheet_includes_partner_name_and_localised_period(tmp_path: Path):
    store, partner = _build_partner(tmp_path, name="Acme  ")
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    assert sheet["B2"].value == "Acme"  # trailing whitespace trimmed
    assert sheet["B3"].value == "Сентябрь 2026"
    assert [sheet.cell(10, column).value for column in range(1, 11)] == list(EXPECTED_HEADERS_RU)


def test_en_sheet_includes_english_period_and_headers(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    assert sheet["A1"].value == "Integration Hub Partner Report"
    assert sheet["B3"].value == "September 2026"
    assert sheet["B2"].value == "Acme Bookings, Inc."
    assert [sheet.cell(10, column).value for column in range(1, 11)] == list(EXPECTED_HEADERS_EN)


def test_report_currency_row_uses_partner_currency(tmp_path: Path):
    store, partner = _build_partner(tmp_path, currency="EUR")
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    assert sheet["A5"].value == "Report currency"
    assert sheet["B5"].value == "EUR"


def test_contact_and_email_rows_are_present(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    assert sheet["A6"].value == "Contact person"
    assert sheet["B6"].value == "Daniel Carter, CFO"
    assert sheet["A7"].value == "Email for questions"
    assert sheet["B7"].value == "billing@acme.io"


def test_rate_row_is_decimal_when_set(tmp_path: Path):
    store, partner = _build_partner(tmp_path, rate=0.2)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    assert sheet["A4"].value == "Ставка Altegio (revenue share)"
    assert sheet["B4"].value == 0.2
    assert sheet["B4"].number_format == "0.00"


def test_rate_row_is_empty_when_partner_has_no_rate(tmp_path: Path):
    store, partner = _build_partner(tmp_path, rate=None)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    assert sheet["B4"].value is None


def test_altegio_rate_column_is_prefilled_in_data_rows(tmp_path: Path):
    """Column I on the primary sheet carries partner.rate by default."""
    store, partner = _build_partner(tmp_path, rate=0.2)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        assert sheet.cell(row_number, 9).value == 0.2


def test_settlement_currency_column_is_prefilled_with_partner_currency(tmp_path: Path):
    """Column G on the primary sheet carries partner.currency."""
    store, partner = _build_partner(tmp_path, currency="EUR")
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        assert sheet.cell(row_number, 7).value == "EUR"


def test_data_rows_are_all_empty_except_rate_and_settlement(tmp_path: Path):
    store, partner = _build_partner(tmp_path, rate=0.2, currency="EUR")
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        for column in range(1, 11):
            if column in (7, 9):
                continue  # G settlement + I rate are prefilled
            assert sheet.cell(row_number, column).value is None, (
                f"row {row_number} column {column} unexpectedly has "
                f"{sheet.cell(row_number, column).value!r}"
            )


def test_data_row_count_is_50(tmp_path: Path):
    assert DATA_ROW_COUNT == 50
    assert LAST_DATA_ROW - FIRST_DATA_ROW + 1 == DATA_ROW_COUNT


def test_total_row_contains_sum_formulas(tmp_path: Path):
    """Ten-column totals: H (Received by partner) and J (Due to Altegio)."""
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    expected_received_total = f"=SUM(H{FIRST_DATA_ROW}:H{LAST_DATA_ROW})"
    expected_due_total = f"=SUM(J{FIRST_DATA_ROW}:J{LAST_DATA_ROW})"
    assert sheet.cell(TOTAL_ROW, 1).value == "ИТОГО в валюте отчёта"
    assert sheet.cell(TOTAL_ROW, 8).value == expected_received_total
    assert sheet.cell(TOTAL_ROW, 10).value == expected_due_total


def test_disclaimer_header_and_body_are_present(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    assert sheet.cell(DISCLAIMER_HEADER_ROW, 1).value == (
        "Ответственность партнёра за достоверность данных"
    )
    body = sheet.cell(DISCLAIMER_BODY_ROW, 1).value
    assert body.startswith("Отправляя этот отчёт, партнёр принимает")


def test_en_sheet_uses_english_disclaimer(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "en")]
    assert sheet.cell(DISCLAIMER_HEADER_ROW, 1).value == (
        "Partner's responsibility for data accuracy"
    )
    body = sheet.cell(DISCLAIMER_BODY_ROW, 1).value
    assert body.startswith("By submitting this report, the partner assumes")


def test_sheet_visual_settings(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="ru")
    workbook = load_workbook(io.BytesIO(blob))
    sheet = workbook[_primary_sheet_title(workbook, "2026-09", "ru")]
    assert sheet.sheet_view.showGridLines is False
    assert sheet.freeze_panes == "A11"
    for letter, width in COLUMN_WIDTHS.items():
        assert sheet.column_dimensions[letter].width == width


def test_example_sheet_contains_filled_rows_with_formulas(tmp_path: Path):
    store, partner = _build_partner(tmp_path)
    blob = build_report_template_xlsx(store, partner, "2026-09", language="en")
    workbook = load_workbook(io.BytesIO(blob), data_only=False)
    example = workbook[TRANSLATIONS["ru"]["example_sheet_name"]]
    # First example row has all the inputs the partner would type plus
    # the prefilled settlement currency (USD) and formulas.
    row_number = FIRST_DATA_ROW
    assert example.cell(row_number, 1).value is not None  # location ID (A)
    assert example.cell(row_number, 2).value  # client name (B)
    assert example.cell(row_number, 3).value  # subscription period (C)
    assert example.cell(row_number, 4).value  # payment currency (D)
    assert isinstance(example.cell(row_number, 5).value, (int, float))  # amount (E)
    assert isinstance(example.cell(row_number, 6).value, (int, float))  # fx (F)
    assert example.cell(row_number, 7).value == "USD"  # G settlement (USD per Kate)
    assert example.cell(row_number, 8).value == f"=E{row_number}*F{row_number}"  # H received
    assert example.cell(row_number, 9).value == 0.2  # I rate
    assert example.cell(row_number, 10).value == f"=H{row_number}*I{row_number}"  # J due