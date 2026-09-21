import base64
import json
import re
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

from revshare.extract import location_ids_from_file, parse_partner_report
from revshare.store import (
    Store,
    assess,
    discrepancy_items,
    normalize_application_links,
    normalize_ids,
    verify_claim,
)
from revshare.web import create_app


def test_normalize_ids():
    assert normalize_ids("101, 102\n103;104\n101") == ["101", "102", "103", "104", "101"]


def test_location_ids_from_file_prefers_the_registry_column(tmp_path: Path):
    path = tmp_path / "report.csv"
    path.write_text(
        "Salon ID;Turnover\n1332009;$1 200,00\n655013;$980,00\n",
        encoding="utf-8",
    )
    ids, source = location_ids_from_file(path)
    assert ids == ["1332009", "655013"]
    assert source == "column"


def test_location_ids_from_file_without_a_registry(tmp_path: Path):
    path = tmp_path / "total.csv"
    path.write_text("Total turnover\n$8 400,00\n", encoding="utf-8")
    ids, source = location_ids_from_file(path)
    assert ids == []
    assert source == "not_found"


def test_parse_partner_report_reads_line_based_totals(tmp_path: Path):
    path = tmp_path / "report.txt"
    path.write_text(
        "Business Name username ID Altegio cost ($) Altegio Reward 20% ($) RATE TOTALS\n"
        "First Beauty Salon 1354377 109,0 21,8\n"
        "Second Studio 84582 79,0 15,8\n"
        "13 BARBERSHOP Centru (Pushkin 44) 1366611 30,0 6,0\n"
        "TOTAL PARTNER $2 173\n"
        "TOTAL ALTEGIO REWARD $435\n",
        encoding="utf-8",
    )
    parsed = parse_partner_report(path, rate=0.20, default_currency="USD")
    assert parsed.location_ids == ["1354377", "84582", "1366611"]
    assert parsed.turnover == 2173
    assert parsed.reported_share == 435
    assert parsed.notes == []


def test_parse_partner_report_calculates_share(tmp_path: Path):
    path = tmp_path / "report.csv"
    path.write_text(
        "Location ID,Turnover\n1001,$120.00\n1002,$80.00\n",
        encoding="utf-8",
    )
    parsed = parse_partner_report(path, rate=0.20, default_currency="USD")
    assert parsed.location_ids == ["1001", "1002"]
    assert parsed.turnover == 200
    assert parsed.reported_share == 40
    assert parsed.currency == "USD"
    assert "share_calculated" in parsed.notes


def test_full_report_round_trip_reads_file_reasons(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
        payment_ids=["101"],
        location_details={"102": {"name": "Demo Salon"}},
    )
    original = tmp_path / "original.csv"
    original.write_text("Location ID,Turnover\n101,$100\n", encoding="utf-8")
    with original.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("original.csv", handle, "text/csv")},
        )
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    full = client.get(f"/submit/demo-partner/review/{draft_token}/full-report.csv")
    assert full.status_code == 200
    assert "Название салона" in full.text
    assert "Demo Salon" in full.text
    assert ",,102,Demo Salon,,USD" in full.text

    completed = tmp_path / "completed.csv"
    completed.write_text(
        "Application ID,Application,Location ID,Название салона,Оборот,Валюта,Статус оплаты,Комментарий\n"
        ",AssistBot,101,,100,USD,Оплата получена,\n"
        ",AssistBot,102,Demo Salon,,USD,Бесплатный месяц,\n"
        "Итого оборот,,,,100,USD,,\n"
        "Итого revshare,,,,20,USD,,\n",
        encoding="utf-8",
    )
    with completed.open("rb") as handle:
        reviewed = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("completed.csv", handle, "text/csv")},
        )
    assert reviewed.status_code == 200
    assert "Show all rows (2)" in reviewed.text
    assert "Free month" in reviewed.text
    reviewed_token = str(reviewed.url).rstrip("/").split("/")[-1]
    confirmed = client.post(f"/submit/demo-partner/confirm/{reviewed_token}", data={"comment": ""})
    assert confirmed.status_code == 200


def test_parse_recommended_partner_report_without_cached_formulas(tmp_path: Path):
    path = tmp_path / "partner-report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report (EN)"
    sheet.append(["Integration Hub Partner Report"])
    sheet.append([])
    sheet.append(
        [
            "Altegio location ID",
            "Client name",
            "Paid subscription period",
            "Payment currency",
            "Received by partner (in payment currency)",
            "Monthly avg rate to USD",
            "Received by partner (USD, report currency)",
            "Altegio rate",
            "Due to Altegio (USD)",
        ]
    )
    sheet.append([312045, "Beauty Loft", "Monthly", "USD", 40, 1, "=E4*F4", 0.2, "=G4*H4"])
    sheet.append([215890, "Atlas", "Monthly", "EUR", 50, 1.1, "=E5*F5", 0.2, "=G5*H5"])
    sheet.append(["TOTAL in report currency (USD)", "", "", "", "", "", "=SUM(G4:G5)", "", "=SUM(I4:I5)"])
    workbook.save(path)

    parsed = parse_partner_report(path, rate=0.2, default_currency="USD")
    assert parsed.location_ids == ["312045", "215890"]
    assert parsed.location_source == "recommended_template"
    assert parsed.turnover == 95
    assert parsed.reported_share == 19
    assert parsed.currency == "USD"
    assert parsed.notes == ["recommended_template"]


def test_parse_monthly_workbook_with_float_ids(tmp_path: Path):
    path = tmp_path / "partner-months.xlsx"
    workbook = Workbook()
    july = workbook.active
    july.title = "Июль 2026"
    august = workbook.create_sheet("Август 2026")
    for sheet, rows, total in (
        (july, [("Salon One", "297134.0", "57.0"), ("Salon Two", "726578.0", "49.0")], 106),
        (august, [("Salon Three", "1295757.0", "80.0")], 80),
    ):
        sheet.append(["Business Name", "ID Altegio", " cost ($)"])
        for name, location_id, cost in rows:
            sheet.append([name, location_id, cost])
        sheet.append([None, None, None])
        sheet.append(["TOTAL AI BeautyBot", None, str(total)])
    workbook.save(path)

    august_report = parse_partner_report(
        path, rate=0.2, default_currency="USD", period="2026-08"
    )
    assert august_report.location_ids == ["1295757"]
    assert august_report.turnover == 80
    assert august_report.reported_share == 16

    july_report = parse_partner_report(
        path, rate=0.2, default_currency="USD", period="2026-07"
    )
    assert july_report.location_ids == ["297134", "726578"]
    assert july_report.turnover == 106

    ids, source = location_ids_from_file(path, "2026-07")
    assert ids == ["297134", "726578"]
    assert source == "column"


def test_normalize_application_links():
    assert normalize_application_links(
        ["https://app.alteg.io/app/123", "https://app.alteg.io/app/456"],
        ["Main app", "Second app"],
    ) == [
        {
            "url": "https://app.alteg.io/app/123",
            "comment": "Main app",
            "category": "",
        },
        {
            "url": "https://app.alteg.io/app/456",
            "comment": "Second app",
            "category": "",
        },
    ]


def test_assessment_finds_missing_and_unknown():
    review = assess(
        ["101", "999"],
        connected_ids={"101", "102", "103"},
        active_ids={"101", "102"},
    )
    assert review.missing_active == ["102"]
    assert review.unknown == ["999"]
    assert review.level in {"yellow", "red"}


def test_discrepancy_items_only_returns_the_delta():
    assert discrepancy_items(
        ["101", "999"],
        connected_ids={"101", "102", "103"},
        active_ids={"101", "102"},
    ) == [
        {"location_id": "102", "kind": "active_not_reported", "required": True},
        {"location_id": "103", "kind": "connected_not_active", "required": False},
        {"location_id": "999", "kind": "reported_not_connected", "required": False},
    ]


def test_payment_callback_contradicts_no_payment_claim():
    assert verify_claim(
        "active_not_reported",
        "no_payment",
        is_connected=True,
        is_active=True,
        had_payment=True,
    ) == ("contradicted", "payment_found")


def test_full_submission_flow(tmp_path: Path):
    db_path = tmp_path / "test.db"
    app = create_app(db_path)
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
    )

    registry = tmp_path / "registry.csv"
    registry.write_text(
        "Location ID,Turnover\n101,$400\n102,$600\n", encoding="utf-8"
    )
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    assert preview.status_code == 200
    assert "Check your report" in preview.text
    assert "1 000 $" in preview.text
    assert 'name="turnover"' not in preview.text
    draft_token = str(preview.url).rstrip("/").split("/")[-1]

    response = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "1000",
            "currency": "USD",
            "reported_share": "300",
            "location_ids": "101\n102",
            "comment": "Monthly registry",
        },
    )
    assert response.status_code == 200
    assert "Saved for reconciliation with Altegio data" in response.text
    dashboard = client.get("/?period=2026-08")
    assert dashboard.status_code == 200
    assert "1 000.00 USD" in dashboard.text
    export = client.get("/export/2026-08.csv")
    assert export.status_code == 200
    assert b"Demo Integration Hub App" in export.content


def test_multi_application_submission_flow(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_application_links(
        partner["id"],
        [
            {
                "url": "https://app.alteg.io/e/mp_39_message_help/",
                "comment": "Message.help",
                "category": "",
            },
            {
                "url": "https://app.alteg.io/e/mp_564_ai_beauty_bot/",
                "comment": "AI Beauty Bot",
                "category": "",
            },
        ],
    )
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101", "102"], active_ids=["101", "102"]
    )
    template = client.get("/submit/demo-partner/template.csv?period=2026-08")
    assert template.status_code == 200
    assert b"Application ID,Application,Location ID,Turnover,Currency,Payment status,Comment" in template.content
    assert b"39,Message.help" in template.content
    assert b"Application URL" not in template.content
    assert b"https://app.alteg.io" not in template.content
    content = (
        "Application,Location ID,Turnover,Currency\n"
        "Message.help,101,100,USD\n"
        "AI Beauty Bot,102,200,USD\n"
    ).encode()

    preview = client.post(
        "/submit/demo-partner/parse",
        data={"period": "2026-08"},
        files={"attachment": ("multi.csv", content, "text/csv")},
    )

    assert preview.status_code == 200
    assert "Message.help" in preview.text
    assert "AI Beauty Bot" in preview.text
    assert "https://app.alteg.io" not in preview.text
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    draft = store.report_draft(draft_token, partner["id"])
    assert [item["application_match"] for item in draft["application_breakdown"]] == [
        "matched",
        "matched",
    ]

    response = client.post(f"/submit/demo-partner/confirm/{draft_token}", data={})

    assert response.status_code == 200
    report_id = store.report_id_for(partner["id"], "2026-08")
    report = store.report(report_id)
    assert [item["application_name"] for item in report["application_breakdown"]] == [
        "Message.help",
        "AI Beauty Bot",
    ]


def test_confirm_ignores_edited_extracted_values(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101"],
        active_ids=["101"],
    )
    registry = tmp_path / "registry.csv"
    registry.write_text("Location ID,Turnover\n101,$400\n", encoding="utf-8")
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    confirmed = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "99999",
            "currency": "EUR",
            "reported_share": "1",
            "location_ids": "999",
            "comment": "",
        },
    )
    assert confirmed.status_code == 200
    report_id = next(
        item["report_id"]
        for item in store.dashboard("2026-08")
        if item["id"] == partner["id"]
    )
    report = store.report(report_id)
    assert report["turnover"] == 400
    assert report["currency"] == "USD"
    assert report["clients"] == ["101"]


def test_registry_is_read_from_the_uploaded_file(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["501", "502"],
        active_ids=["501", "502"],
    )
    registry = tmp_path / "registry.csv"
    registry.write_text(
        "Location ID,Client,Amount\n501,Salon One,$120.00\n502,Salon Two,$80.00\n",
        encoding="utf-8",
    )

    with registry.open("rb") as handle:
        response = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )

    assert response.status_code == 200
    assert "Salons</span>" in response.text
    assert ">2<" in response.text
    assert "200 $" in response.text
    assert "60 $" in response.text
    assert 'name="turnover"' not in response.text
    assert store.dashboard("2026-08")[0]["report_id"] is None

    draft_token = str(response.url).rstrip("/").split("/")[-1]
    confirmed = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "200",
            "currency": "USD",
            "reported_share": "60",
            "location_ids": "501\n502",
            "comment": "",
        },
    )
    assert confirmed.status_code == 200
    assert "Saved for reconciliation with Altegio data" in confirmed.text
    report_id = next(
        item["report_id"]
        for item in store.dashboard("2026-08")
        if item["id"] == partner["id"]
    )
    assert store.report(report_id)["clients"] == ["501", "502"]


def test_review_page_rereads_the_stored_file(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    registry = tmp_path / "registry.csv"
    registry.write_text(
        "Location ID,Amount\n701,$50\n702,$70\n", encoding="utf-8"
    )
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    draft_token = str(preview.url).rstrip("/").split("/")[-1]

    store.update_report_draft(
        token=draft_token,
        partner_id=partner["id"],
        turnover=999999,
        currency="USD",
        reported_share=1,
        location_ids=["701"],
        extraction_notes=[],
    )

    refreshed = client.get(f"/submit/demo-partner/review/{draft_token}")
    assert refreshed.status_code == 200
    assert 'data-panel="summary"' not in refreshed.text
    assert "Your report" in refreshed.text
    assert "Show all rows (2)" in refreshed.text
    assert "Check your report" in refreshed.text
    assert "120 $" in refreshed.text
    assert "999999" not in refreshed.text
    draft = store.report_draft(draft_token, partner["id"])
    assert draft["location_ids"] == ["701", "702"]


def test_report_without_registry_requires_review(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    report = tmp_path / "total.csv"
    report.write_text("Total turnover,1000 USD\n", encoding="utf-8")
    with report.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("total.csv", handle, "text/csv")},
        )
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    response = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "1000",
            "currency": "USD",
            "reported_share": "300",
            "location_ids": "",
            "comment": "",
        },
    )
    assert response.status_code == 200
    assert "Saved for reconciliation with Altegio data" in response.text
    assert "Material discrepancy found" not in response.text


def test_partner_must_explain_active_gaps_and_admin_sees_evidence(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
        payment_ids=["102"],
        location_details={
            "102": {
                "name": "Demo Salon",
                "connected_at": "2024-03-10",
                "last_api_at": "2026-08-15T12:00:00",
                "has_api_amounts": True,
            }
        },
    )
    registry = tmp_path / "registry.csv"
    registry.write_text("Location ID,Turnover\n101,$100\n", encoding="utf-8")
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    assert ">102<" in preview.text
    assert "Fix 1 rows" in preview.text
    assert "Missing from your file" in preview.text
    assert ">1</strong>" in preview.text
    assert ">2</strong>" in preview.text
    assert 'data-open-tab="differences"' not in preview.text
    assert 'class="card review-form tab-panel active" data-panel="summary"' not in preview.text
    assert "Upload the corrected file" in preview.text
    assert "Contract rate" not in preview.text
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    differences = client.get(
        f"/submit/demo-partner/review/{draft_token}/differences.csv"
    )
    assert differences.status_code == 200
    assert "altegio-active-2026-08.csv" in differences.headers["content-disposition"]
    assert (
        "102,2026-08,Active during reporting month,Missing from partner report,"
        "Demo Salon,2024-03-10,2026-08-15T12:00:00,Yes"
    ) in differences.text
    assert "101" not in differences.text
    assert "Download full report" in preview.text
    assert "Upload filled report" in preview.text

    missing_reason = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "100",
            "currency": "USD",
            "reported_share": "30",
            "location_ids": "101",
            "comment": "",
        },
    )
    assert missing_reason.status_code == 422
    assert "Complete the full report" in missing_reason.text

    confirmed = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "100",
            "currency": "USD",
            "reported_share": "30",
            "location_ids": "101",
            "comment": "",
            "claim_location": "102",
            "claim_kind": "active_not_reported",
            "claim_reason": "no_payment",
            "claim_detail": "",
        },
    )
    assert confirmed.status_code == 200
    assert "Saved for reconciliation with Altegio data" in confirmed.text
    assert "/100" not in confirmed.text
    report_id = next(
        item["report_id"]
        for item in store.dashboard("2026-08")
        if item["id"] == partner["id"]
    )
    admin = client.get(f"/reports/{report_id}")
    assert "Contradicted by data" in admin.text
    assert "A payment callback exists for this month." in admin.text

    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
        payment_ids=[],
    )
    refreshed_admin = client.get(f"/reports/{report_id}")
    assert "Needs review" in refreshed_admin.text
    assert "There is API activity but no payment callback." in refreshed_admin.text


def test_dashboard_filters_and_revshare_column(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    quiet = store.add_partner(
        name="Quiet App",
        legal_name="",
        legal_country="",
        altegio_legal_entity="",
        contact_name="",
        email="quiet@example.com",
        telegram="",
        report_channel="email",
        token="quiet-token",
        rate=0.1,
        currency="USD",
    )
    assert quiet
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=1000,
        currency="USD",
        reported_share=300,
        location_ids=["101"],
        comment="",
        attachment=None,
    )

    page = client.get("/?period=2026-08")
    assert page.status_code == 200
    assert "Revshare" in page.text
    assert "300.00 USD" in page.text
    assert "Shown: 2 of 2" in page.text
    assert "Quiet App" in page.text

    submitted = client.get("/?period=2026-08&status=submitted")
    assert "Quiet App" not in submitted.text
    assert "Shown: 1 of 2" in submitted.text
    assert "Total" in submitted.text
    assert "reports: 1" in submitted.text

    sorted_share = client.get("/?period=2026-08&sort=share&dir=desc")
    assert sorted_share.status_code == 200
    demo_pos = sorted_share.text.index("Demo Integration Hub App")
    quiet_pos = sorted_share.text.index("Quiet App")
    assert demo_pos < quiet_pos

    missing = client.get("/?period=2026-08&status=missing")
    assert "Quiet App" in missing.text
    assert "Demo Integration Hub App" not in missing.text

    searched = client.get("/?period=2026-08&q=quiet@example.com")
    assert "Quiet App" in searched.text
    assert "Shown: 1 of 2" in searched.text

    empty = client.get("/?period=2026-08&status=submitted&risk=red")
    assert "No partners match these filters" in empty.text

    export = client.get("/export/2026-08.csv?status=missing")
    assert export.status_code == 200
    assert b"Demo Integration Hub App" not in export.content


def test_partner_profile_lists_previous_reports(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")

    empty = client.get(f"/admin/partners/{partner['id']}/edit")
    assert "Partner profile" in empty.text
    assert "Reported months" in empty.text
    assert "No reports yet" in empty.text

    for period, turnover in (("2026-07", 500), ("2026-08", 900)):
        store.replace_usage(
            partner["id"], period, connected_ids=["101"], active_ids=["101"]
        )
        store.submit_report(
            partner_id=partner["id"],
            period=period,
            turnover=turnover,
            currency="USD",
            reported_share=turnover * 0.3,
            location_ids=["101"],
            comment="",
            attachment=None,
        )

    history = store.partner_history(partner["id"])
    assert [row["period"] for row in history] == ["2026-08", "2026-07"]
    assert history[0]["expected_share"] == 270.0
    assert history[0]["active_count"] == 1

    page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert "August 2026" in page.text
    assert "July 2026" in page.text
    assert "270.00 USD" in page.text
    assert "history-bar-fill share" in page.text
    assert "history-bar-fill connected" in page.text


def test_partner_profile_splits_into_tabs(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")

    page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert 'data-tab="main"' in page.text
    assert 'data-tab="settings"' in page.text
    assert 'data-tab="history"' in page.text
    assert 'class="card wide-form tab-panel active" data-panel="main"' in page.text

    history_tab = client.get(f"/admin/partners/{partner['id']}/edit?tab=history")
    assert 'data-panel="history" id="history"' in history_tab.text
    assert 'tab-panel active" data-panel="history"' in history_tab.text
    assert 'tab-panel " data-panel="main"' in history_tab.text

    saved = client.post(
        f"/admin/partners/{partner['id']}/notifications",
        data={
            "notice_language": "ru",
            "notice_on": ["invite", "last"],
            "notice_key": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_day": ["3", "4", "9", "10", "11"],
            "own_schedule": "on",
        },
    )
    assert saved.status_code == 200
    assert "Changes saved." in saved.text
    settings = store.partner(partner["id"])["notice_settings"]
    by_key = {item["key"]: item for item in settings["notices"]}
    assert settings["language"] == "ru"
    assert by_key["invite"]["day"] == 3
    assert by_key["reminder"]["enabled"] is False
    assert store.partner(partner["id"])["name"] == partner["name"]


def test_dashboard_shows_verified_and_review_revshare_sums(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=1000,
        currency="USD",
        reported_share=300,
        location_ids=["101"],
        comment="",
        attachment=None,
    )
    clean = store.dashboard("2026-08")[0]
    assert clean["status"] == "verified"

    from revshare.web import report_is_approved, revshare_group

    reported = [item for item in store.dashboard("2026-08") if item["report_id"]]
    assert revshare_group([item for item in reported if report_is_approved(item)])[
        "amount"
    ] == 0
    assert revshare_group([item for item in reported if not report_is_approved(item)])[
        "amount"
    ] == 300

    client.post(f"/reports/{clean['report_id']}/approve", follow_redirects=False)
    reported = [item for item in store.dashboard("2026-08") if item["report_id"]]
    approved_group = revshare_group(
        [item for item in reported if report_is_approved(item)], in_usd=True
    )
    assert approved_group["amount"] == 300
    assert approved_group["currency"] == "USD"
    assert revshare_group([item for item in reported if not report_is_approved(item)])[
        "amount"
    ] == 0


def test_approved_revshare_summary_uses_usd_for_foreign_currency(tmp_path: Path):
    store = Store(tmp_path / "test.db")
    partner_id = store.add_partner(
        name="KZT Partner",
        legal_name="",
        legal_country="",
        altegio_legal_entity="",
        contact_name="",
        email="",
        telegram="",
        report_channel="email",
        token="kzt-partner",
        rate=0.2,
        currency="USD",
    )
    store.replace_usage(
        partner_id, "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    report_id, _ = store.submit_report(
        partner_id=partner_id,
        period="2026-08",
        turnover=19800,
        currency="KZT",
        reported_share=3960,
        location_ids=["101"],
        comment="",
        attachment=None,
    )
    store.approve_report(report_id)
    item = next(row for row in store.dashboard("2026-08") if row["report_id"] == report_id)
    from revshare.web import report_is_approved, revshare_group

    group = revshare_group([item for item in [item] if report_is_approved(item)], in_usd=True)
    assert group["currency"] == "USD"
    assert group["amount"] == item["share_usd"]
    assert group["amount"] != item["expected_share"]


def test_financial_mismatch_requires_review_and_can_be_approved(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    report_id, review = store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=1000,
        currency="USD",
        reported_share=100,
        location_ids=["101"],
        comment="",
        attachment="first.xlsx",
    )
    assert review.level == "yellow"
    assert any(reason["key"] == "reason_share_mismatch" for reason in review.reasons)
    report = store.report(report_id)
    assert report["workflow_status"] == "needs_review"
    detail = client.get(f"/reports/{report_id}")
    assert "Reported revshare is 100" in detail.text
    approved = client.post(f"/reports/{report_id}/approve", follow_redirects=False)
    assert approved.status_code == 303
    assert store.report(report_id)["workflow_status"] == "approved"
    detail = client.get(f"/reports/{report_id}")
    assert "Send to Google Sheet" in detail.text
    report = store.report(report_id)
    notes = [event["note"] for event in report["status_log"]]
    assert "approved" in notes
    assert "submitted" in notes
    detail = client.get(f"/reports/{report_id}")
    assert "Report approved" in detail.text
    from revshare.web import report_is_approved, revshare_group

    reported = [item for item in store.dashboard("2026-08") if item["report_id"]]
    verified = revshare_group([item for item in reported if report_is_approved(item)])
    assert verified["count"] == 1
    assert verified["amount"] == reported[0]["expected_share"]


def test_repeated_submission_keeps_versions_and_resets_approval(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    report_id, _ = store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=100,
        currency="USD",
        reported_share=30,
        location_ids=["101"],
        comment="first",
        attachment="first.xlsx",
    )
    assert store.approve_report(report_id)
    second_id, _ = store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=200,
        currency="USD",
        reported_share=60,
        location_ids=["101"],
        comment="second",
        attachment="second.xlsx",
    )
    assert second_id == report_id
    report = store.report(report_id)
    assert report["workflow_status"] == "auto_checked"
    assert [version["version"] for version in report["versions"]] == [2, 1]
    assert [version["turnover"] for version in report["versions"]] == [200, 100]


def test_same_location_reported_by_two_partners_flags_both(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    store: Store = app.state.store
    first = store.partner_by_token("demo-partner")
    second_id = store.add_partner(
        name="Second App",
        legal_name="",
        legal_country="",
        altegio_legal_entity="",
        contact_name="",
        email="second@example.com",
        telegram="",
        report_channel="email",
        token="second-token",
        rate=0.3,
        currency="USD",
    )
    for partner_id in (first["id"], second_id):
        store.replace_usage(
            partner_id, "2026-08", connected_ids=["101"], active_ids=["101"]
        )
        store.submit_report(
            partner_id=partner_id,
            period="2026-08",
            turnover=100,
            currency="USD",
            reported_share=30,
            location_ids=["101"],
            comment="",
            attachment=None,
        )
    reports = [item for item in store.dashboard("2026-08") if item["report_id"]]
    assert all(item["status"] == "requires_review" for item in reports)
    assert all(
        any(reason["key"] == "reason_cross_partner" for reason in item["reasons"])
        for item in reports
    )


def test_disconnect_registry_requires_consecutive_idle_months(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-06",
        connected_ids=["101", "102", "103"],
        active_ids=[],
    )
    store.replace_usage(
        partner["id"],
        "2026-07",
        connected_ids=["101", "102", "103", "104"],
        active_ids=["102"],
        payment_ids=["103"],
    )
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102", "103", "104"],
        active_ids=[],
    )

    candidates = store.disconnect_candidates("2026-08")
    assert [(item["location_id"], item["months"], item["since"]) for item in candidates] == [
        ("101", 3, "2026-06")
    ]

    store.set_idle_months_threshold("2")
    candidates = store.disconnect_candidates("2026-08")
    assert [(item["location_id"], item["months"]) for item in candidates] == [
        ("101", 3),
        ("104", 2),
    ]

    page = client.get("/admin/usage?period=2026-08")
    assert page.status_code == 200
    assert "Disconnect candidates" in page.text
    assert "101" in page.text
    assert "104" in page.text
    assert 'href="/admin/disconnect-candidates"' not in page.text
    redirected = client.get(
        "/admin/disconnect-candidates?period=2026-08", follow_redirects=False
    )
    assert redirected.status_code == 303
    assert redirected.headers["location"].endswith(
        "/admin/usage?period=2026-08#disconnect-candidates"
    )
    exported = client.get("/admin/disconnect-candidates/2026-08.csv")
    assert exported.status_code == 200
    assert "Demo Integration Hub App,101,3,2026-06,2026-08" in exported.text


def test_partner_sees_long_idle_installations_on_review(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    for period in ("2026-06", "2026-08"):
        store.replace_usage(
            partner["id"],
            period,
            connected_ids=["101", "301", "302"],
            active_ids=["101"],
        )
    store.replace_usage(
        partner["id"],
        "2026-07",
        connected_ids=["101", "301", "302"],
        active_ids=["101", "302"],
    )

    registry = tmp_path / "registry.csv"
    registry.write_text("Location ID,Turnover\n101,$100\n", encoding="utf-8")
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    assert preview.status_code == 200
    assert "Connected but inactive" in preview.text

    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    exported = client.get(
        f"/submit/demo-partner/review/{draft_token}/differences.csv?kind=disconnect"
    )
    assert exported.status_code == 200
    assert "301,2026-08," in exported.text
    assert "Idle since 2026-06, 3 months in a row" in exported.text
    assert "302" not in exported.text


def test_altegio_data_page_links_to_settings(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")

    page = client.get("/admin/usage")
    assert page.status_code == 200
    assert 'href="/admin/settings#altegio"' in page.text
    assert "cannot be verified" in page.text

    settings = client.get("/admin/settings")
    assert settings.status_code == 200
    assert "How a report is checked" in settings.text
    assert 'action="/admin/settings/altegio/import"' in settings.text
    assert 'action="/admin/settings/altegio/sync"' in settings.text

    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101"],
        payment_ids=["101"],
    )
    filled = client.get(f"/admin/usage?partner_id={partner['id']}&period=2026-08")
    assert "installed 2, with activity 1, with a payment callback 1" in filled.text
    assert "cannot be verified" not in filled.text

    other_id = store.add_partner(
        name="Other Test Partner",
        legal_name="",
        legal_country="",
        altegio_legal_entity="",
        contact_name="",
        email="other@example.com",
        telegram="",
        report_channel="email",
        token="other-test-partner",
        rate=0.2,
        currency="USD",
    )
    store.replace_usage(
        other_id,
        "2026-08",
        connected_ids=["901"],
        active_ids=["902"],
        payment_ids=["902"],
    )
    switched = client.get(f"/admin/usage?partner_id={other_id}&period=2026-08")
    assert re.search(r'<textarea id="connected_ids"[^>]*>\s*901', switched.text)
    assert re.search(r'<textarea id="active_ids"[^>]*>\s*902', switched.text)
    assert 'name="partner_id" value="%s"' % other_id in switched.text
    assert 'id="partner_id" name="partner_id" data-auto-submit' in switched.text


def test_dashboard_sorts_by_any_column(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    small = store.add_partner(
        name="Alpha App",
        legal_name="",
        legal_country="",
        altegio_legal_entity="",
        contact_name="",
        email="alpha@example.com",
        telegram="",
        report_channel="email",
        token="alpha-token",
        rate=0.2,
        currency="USD",
    )
    store.replace_usage(small, "2026-08", connected_ids=["301"], active_ids=["301"])
    store.submit_report(
        partner_id=small,
        period="2026-08",
        turnover=100,
        currency="USD",
        reported_share=20,
        location_ids=["301"],
        comment="",
        attachment=None,
    )
    big = store.partner_by_token("demo-partner")
    store.replace_usage(
        big["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    store.submit_report(
        partner_id=big["id"],
        period="2026-08",
        turnover=5000,
        currency="USD",
        reported_share=1500,
        location_ids=["101"],
        comment="",
        attachment=None,
    )

    by_name = client.get("/?period=2026-08").text
    assert by_name.index("Alpha App") < by_name.index("Demo Integration Hub App")

    desc = client.get("/?period=2026-08&sort=turnover&dir=desc").text
    assert desc.index("Demo Integration Hub App") < desc.index("Alpha App")

    asc = client.get("/?period=2026-08&sort=turnover&dir=asc").text
    assert asc.index("Alpha App") < asc.index("Demo Integration Hub App")

    share_desc = client.get("/?period=2026-08&sort=share&dir=desc").text
    assert share_desc.index("Demo Integration Hub App") < share_desc.index("Alpha App")
    assert "sort=turnover&amp;dir=desc" in by_name


def test_non_usd_reports_also_show_dollars(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.update_partner(
        partner_id=partner["id"],
        name=partner["name"],
        legal_name="",
        legal_country="",
        altegio_legal_entity="Open HC",
        contact_name="",
        email="demo@example.com",
        telegram="",
        report_channel="email",
        rate=0.2,
        currency="EUR",
    )
    store.replace_usage(
        partner["id"], "2026-08", connected_ids=["101"], active_ids=["101"]
    )
    report_id, _ = store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=920,
        currency="EUR",
        reported_share=184,
        location_ids=["101"],
        comment="",
        attachment=None,
    )

    page = client.get("/?period=2026-08")
    assert "920.00 EUR" in page.text
    assert "1 000.00 USD" in page.text
    assert "184.00 EUR" in page.text
    assert "200.00 USD" in page.text

    detail = client.get(f"/reports/{report_id}")
    assert "1 000.00 USD" in detail.text

    rates = client.post("/admin/rates", data={"rate_EUR": "0.5"})
    assert rates.status_code == 200
    assert store.fx_rates()["EUR"] == 0.5
    updated = client.get("/?period=2026-08")
    assert "1 840.00 USD" in updated.text


def test_settings_load_the_month_average_on_refresh(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    app.state.store.fx_loader = lambda period: {
        "EUR": 0.77,
        "USD": 1.0,
        "_days": 30,
    }
    client = TestClient(app)
    page = client.get("/admin/settings?period=2026-08")
    assert page.status_code == 200
    assert "Month these rates belong to" in page.text
    assert "Load month average" in page.text
    assert "0.77" in page.text
    assert "Average of daily rates for 2026-08, 30 days." in page.text


def test_admin_can_preview_the_report_email(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")

    preview = client.get(f"/admin/partners/{partner['id']}/letter/invite")
    assert preview.status_code == 200
    assert "Altegio Integration Hub" in preview.text
    assert "Open the report form" in preview.text
    assert "/submit/demo-partner?period=" in preview.text

    card = client.get(f"/admin/partners/{partner['id']}/edit")
    assert f"/admin/partners/{partner['id']}/letter/invite" in card.text
    assert client.get(f"/admin/partners/{partner['id']}/letter/digest").status_code == 404


def test_partner_can_submit_without_explaining_connected_inactive(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102", "201", "202", "203"],
        active_ids=["101", "102"],
    )
    registry = tmp_path / "registry.csv"
    registry.write_text("Location ID,Turnover\n101,$100\n102,$80\n", encoding="utf-8")
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    assert preview.status_code == 200
    assert "Location ID 201" not in preview.text
    assert "Connected but inactive" in preview.text
    assert "Ready to submit" in preview.text
    assert "Explain active clients missing from the report" not in preview.text
    draft_token = str(preview.url).rstrip("/").split("/")[-1]

    idle_list = client.get(
        f"/submit/demo-partner/review/{draft_token}/differences.csv?kind=connected"
    )
    assert idle_list.status_code == 200
    assert '201,2026-08,"Connected at month end, no API activity"' in idle_list.text
    assert "101" not in idle_list.text

    confirmed = client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "180",
            "currency": "USD",
            "reported_share": "54",
            "location_ids": "101\n102",
            "comment": "",
            "idle_declaration": "no_payment",
        },
    )
    assert confirmed.status_code == 200
    assert "Saved for reconciliation with Altegio data" in confirmed.text

    report_id = store.report_id_for(partner["id"], "2026-08")
    saved = store.report(report_id)
    assert saved["claims"] == []
    assert saved["idle_claims"] == {"total": 3, "contradicted": []}

    admin = client.get(f"/reports/{report_id}")
    assert "Partner declared no payment" in admin.text
    assert "these installations can be disconnected" in admin.text

    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102", "201", "202", "203"],
        active_ids=["101", "102"],
        payment_ids=["202"],
    )
    rechecked = client.get(f"/reports/{report_id}")
    assert "Payment callbacks contradict the declaration, locations: 1" in rechecked.text
    assert store.report(report_id)["idle_claims"]["contradicted"] == ["202"]


def test_next_report_resolves_payment_follow_up(tmp_path: Path):
    store = Store(tmp_path / "test.db")
    partner = store.partner_by_token("demo-partner")
    store.replace_usage(
        partner["id"],
        "2026-07",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
    )
    july_id, _ = store.submit_report(
        partner_id=partner["id"],
        period="2026-07",
        turnover=100,
        currency="USD",
        reported_share=30,
        location_ids=["101"],
        comment="",
        attachment=None,
        claims=[
            {
                "location_id": "102",
                "kind": "active_not_reported",
                "reason": "pay_next_month",
                "detail": "",
            }
        ],
    )
    assert store.report(july_id)["claims"][0]["verification_status"] == "follow_up"

    store.replace_usage(
        partner["id"],
        "2026-08",
        connected_ids=["101", "102"],
        active_ids=["101", "102"],
    )
    store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=200,
        currency="USD",
        reported_share=60,
        location_ids=["101", "102"],
        comment="",
        attachment=None,
    )
    july_claim = store.report(july_id)["claims"][0]
    assert july_claim["verification_status"] == "confirmed"
    assert july_claim["verification_evidence"] == "appeared_next_report"


def test_admin_switcher_offers_english_and_russian(tmp_path: Path):
    client = TestClient(create_app(tmp_path / "test.db"))
    response = client.get("/")
    assert "Integration Hub Revshare" in response.text
    assert "Overview" in response.text
    assert "<select data-language" in response.text
    assert "Русский" in response.text
    assert "Українська" not in response.text
    assert "Polski" not in response.text
    assert "Português (Brasil)" not in response.text
    assert "Magyar" not in response.text
    assert response.text.index("English") < response.text.index("Русский")
    assert "Хаб Интеграций" not in response.text

    form = client.get("/submit/demo-partner")
    assert "Magyar" in form.text
    assert "Português (Brasil)" in form.text

    switched = client.get("/language/pl?next=/")
    assert switched.status_code == 200
    assert "Sierpień" in switched.text

    switched = client.get("/language/uk?next=/admin/partners")
    assert switched.status_code == 200
    assert "Хаб Інтеграцій" in switched.text
    assert "Назва юридичної особи" in switched.text
    assert "Швидко додати" in switched.text
    edit = client.get("/admin/partners/1/edit")
    assert "Україна" in edit.text
    assert "Albania" not in edit.text


def test_admin_auth_protects_internal_routes_but_not_partner_form(
    tmp_path: Path, monkeypatch
):
    monkeypatch.setenv("REVSHARE_ADMIN_USER", "kate")
    monkeypatch.setenv("REVSHARE_ADMIN_PASSWORD", "secret")
    client = TestClient(create_app(tmp_path / "test.db", admin_auth=True))
    denied = client.get("/")
    assert denied.status_code == 401
    assert denied.headers["www-authenticate"].startswith("Basic")
    assert client.get("/submit/demo-partner").status_code == 200
    credentials = base64.b64encode(b"kate:secret").decode()
    allowed = client.get("/", headers={"Authorization": f"Basic {credentials}"})
    assert allowed.status_code == 200


def test_admin_auth_fails_closed_without_password(tmp_path: Path, monkeypatch):
    monkeypatch.delenv("REVSHARE_ADMIN_PASSWORD", raising=False)
    client = TestClient(create_app(tmp_path / "test.db", admin_auth=True))

    assert client.get("/").status_code == 503
    assert client.get("/submit/demo-partner").status_code == 200


def test_admin_auth_can_be_disabled_explicitly_for_local_preview(monkeypatch):
    monkeypatch.setenv("REVSHARE_ADMIN_AUTH", "0")
    monkeypatch.delenv("REVSHARE_ADMIN_PASSWORD", raising=False)
    client = TestClient(create_app())

    assert client.get("/").status_code == 200


def test_admin_posts_require_csrf_token(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("REVSHARE_ADMIN_USER", "kate")
    monkeypatch.setenv("REVSHARE_ADMIN_PASSWORD", "secret")
    app = create_app(tmp_path / "test.db", admin_auth=True)
    client = TestClient(app)
    credentials = base64.b64encode(b"kate:secret").decode()
    headers = {"Authorization": f"Basic {credentials}"}

    page = client.get("/admin/settings", headers=headers)
    assert page.status_code == 200
    token = re.search(r'name="csrf_token" value="([^"]+)"', page.text).group(1)

    denied = client.post("/admin/settings/telegram", headers=headers, data={})
    assert denied.status_code == 403
    allowed = client.post(
        "/admin/settings/telegram",
        headers=headers,
        data={"csrf_token": token, "digest_chat": ""},
        follow_redirects=False,
    )
    assert allowed.status_code == 303


def test_accrual_page_does_not_write_to_sheets_on_get(tmp_path: Path):
    client = TestClient(create_app(tmp_path / "test.db"))
    with patch("revshare.web.flush_queued_sheet_writes") as flush:
        response = client.get("/admin/accrual")
    assert response.status_code == 200
    flush.assert_not_called()


def test_cross_origin_posts_are_rejected(tmp_path: Path):
    client = TestClient(create_app(tmp_path / "test.db"))
    response = client.post(
        "/admin/notices/run",
        headers={"Origin": "https://attacker.example"},
    )
    assert response.status_code == 403


def test_partner_form_opens_in_the_language_kate_picked(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    store.set_partner_notifications(partner["id"], {"language": "hu"})

    form = client.get("/submit/demo-partner")
    assert "Havi jelentés" in form.text
    assert "Jelentési hónap" in form.text
    assert "Reporting month" not in form.text
    assert 'href="/admin/partners"' not in form.text
    assert 'href="/admin/usage"' not in form.text

    client.get("/language/pt?next=/submit/demo-partner")
    chosen = client.get("/submit/demo-partner")
    assert "Relatório mensal" in chosen.text
    assert "Havi jelentés" not in chosen.text


def test_russian_country_names_and_canonical_storage(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    client.get("/language/ru?next=/")
    client.post(
        "/admin/partners",
        data={"name": "RU App", "legal_name": "ООО Тест", "rate": "20"},
    )
    partner = next(row for row in app.state.store.partners() if row["name"] == "RU App")
    edit_page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert "Албания" in edit_page.text
    assert "ОАЭ" in edit_page.text
    assert "Albania" not in edit_page.text

    client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "RU App",
            "legal_name": "ООО Тест",
            "legal_country": "Украина",
            "altegio_legal_entity": "Альт тех",
            "contact_name": "Иван",
            "email": "ivan@example.com",
            "telegram": "",
            "report_channel": "email",
            "application_url": "https://app.alteg.io/app/7",
            "application_comment": "",
            "rate": "20",
            "currency": "RUB",
        },
    )
    partner = app.state.store.partner(partner["id"])
    assert partner["legal_country"] == "Ukraine"
    edit_page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert 'value="Украина"' in edit_page.text
    assert "Напоминания и даты" in edit_page.text


def test_partner_legal_entity_fields_are_saved(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    response = client.post(
        "/admin/partners",
        data={
            "name": "Ukrainian App",
            "legal_name": "ТОВ Інтеграція",
            "rate": "25",
        },
    )
    assert response.status_code == 200
    partner = next(row for row in app.state.store.partners() if row["name"] == "Ukrainian App")
    assert partner["legal_name"] == "ТОВ Інтеграція"
    assert partner["rate"] == 0.25

    edit_page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert edit_page.status_code == 200
    assert "Open HC" in edit_page.text
    assert "Altegio KFT" in edit_page.text
    assert "Altegio LTD" in edit_page.text
    assert "Альт тех" in edit_page.text
    assert "United Arab Emirates" in edit_page.text
    assert "Chat bots" in edit_page.text
    assert "SMS providers" in edit_page.text
    assert "Online booking and widgets" not in edit_page.text

    updated = client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Ukrainian App Updated",
            "legal_name": "ТОВ Нова Інтеграція",
            "legal_country": "Poland",
            "altegio_legal_entity": "Open HC",
            "contact_name": "Anna",
            "email": "finance@example.com",
            "telegram": "@team",
            "report_channel": "email",
            "internal_comment": "Проверить новый договор в октябре",
            "application_url": [
                "https://app.alteg.io/app/42",
                "https://app.alteg.io/app/43",
            ],
            "application_comment": ["Main", "Extra"],
            "application_category": ["chat_bot", "sms"],
            "rate": "30",
            "currency": "EUR",
        },
    )
    assert updated.status_code == 200
    saved = app.state.store.partner(partner["id"])
    assert saved["name"] == "Ukrainian App Updated"
    assert saved["legal_country"] == "Poland"
    assert saved["altegio_legal_entity"] == "Open HC"
    assert saved["contact_name"] == "Anna"
    assert saved["report_channel"] == "email"
    assert saved["internal_comment"] == "Проверить новый договор в октябре"
    assert saved["rate"] == 0.30
    assert app.state.store.application_links(partner["id"]) == [
        {
            "url": "https://app.alteg.io/app/42",
            "comment": "Main",
            "category": "chat_bot",
        },
        {
            "url": "https://app.alteg.io/app/43",
            "comment": "Extra",
            "category": "sms",
        },
    ]
    partner_form = client.get(f"/submit/{saved['token']}")
    assert "https://app.alteg.io/app/42" in partner_form.text
    assert "https://app.alteg.io/app/43" in partner_form.text
    edit_page = client.get(f"/admin/partners/{partner['id']}/edit")
    assert 'href="https://app.alteg.io/app/42"' in edit_page.text
    assert f'href="/submit/{saved["token"]}?period=' in edit_page.text
    assert "Проверить новый договор в октябре" in edit_page.text
    assert "Delete partner" in client.get(f"/admin/partners/{partner['id']}/edit").text
    assert "Delete partner" not in client.get("/admin/partners").text

    deleted = client.post(f"/admin/partners/{partner['id']}/delete")
    assert deleted.status_code == 200
    assert all(row["id"] != partner["id"] for row in app.state.store.partners())
    assert client.get(f"/admin/partners/{partner['id']}/edit").status_code == 404
    assert client.get(f"/submit/{saved['token']}").status_code == 404


def test_partner_can_be_saved_without_altegio_legal_entity(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    client.post(
        "/admin/partners",
        data={"name": "Partner Without Entity", "legal_name": "Partner LLC", "rate": "15"},
    )
    partner = next(
        row for row in app.state.store.partners() if row["name"] == "Partner Without Entity"
    )

    response = client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Partner Without Entity Updated",
            "legal_name": "Partner LLC",
            "legal_country": "",
            "report_channel": "email",
            "email": "partner@example.com",
            "telegram": "",
            "application_url": "https://app.alteg.io/e/mp_38_assist/",
            "application_comment": "",
            "rate": "15",
            "currency": "USD",
        },
    )

    assert response.status_code == 200
    saved = app.state.store.partner(partner["id"])
    assert saved["name"] == "Partner Without Entity Updated"
    assert saved["altegio_legal_entity"] == ""


def test_partner_can_be_created_and_saved_without_application_link(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)

    response = client.post("/admin/partners", data={"name": "", "legal_name": ""})
    assert response.status_code == 200
    partner = app.state.store.partners()[-1]
    assert partner["name"].startswith("Новый партнёр ")

    response = client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Партнёр без приложения",
            "legal_name": "",
            "legal_country": "",
            "altegio_legal_entity": "",
            "contact_name": "",
            "email": "",
            "telegram": "",
            "report_channel": "email",
            "application_url": "",
            "application_comment": "",
            "rate": "",
            "currency": "USD",
        },
    )
    assert response.status_code == 200
    saved = app.state.store.partner(partner["id"])
    assert saved["name"] == "Партнёр без приложения"
    assert app.state.store.application_links(partner["id"]) == []


def test_partner_can_be_saved_without_notification_contact(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    client.post(
        "/admin/partners",
        data={"name": "Partner Without Contact", "legal_name": "Partner LLC", "rate": "15"},
    )
    partner = next(
        row for row in app.state.store.partners() if row["name"] == "Partner Without Contact"
    )

    response = client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Partner Without Contact Updated",
            "legal_name": "Partner LLC",
            "report_channel": "email",
            "application_url": "https://app.alteg.io/e/mp_38_assist/",
            "rate": "15",
            "currency": "USD",
        },
    )

    assert response.status_code == 200
    assert app.state.store.partner(partner["id"])["name"] == "Partner Without Contact Updated"


def test_partner_notice_exceptions_are_saved(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    client.post(
        "/admin/partners",
        data={"name": "Late Reporter", "legal_name": "Late LLC", "rate": "15"},
    )
    partner = next(row for row in app.state.store.partners() if row["name"] == "Late Reporter")
    edit = client.get(f"/admin/partners/{partner['id']}/edit")
    assert "Reminders and dates" in edit.text
    assert "Saturday and Sunday dates move to Monday automatically." in edit.text
    assert "Move Saturday and Sunday dates to Monday" not in edit.text
    assert "Message text" not in edit.text
    assert "/admin/settings?language=en#notices" in edit.text
    assert "/admin/settings#schedule" in edit.text
    assert "Use personal dates for this partner" in edit.text
    saved = client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Late Reporter",
            "legal_name": "Late LLC",
            "legal_country": "Estonia",
            "altegio_legal_entity": "Altegio KFT",
            "contact_name": "Eva",
            "email": "eva@example.com",
            "telegram": "",
            "report_channel": "email",
            "application_url": "https://app.alteg.io/app/99",
            "application_comment": "",
            "rate": "15",
            "currency": "EUR",
            "notice_on": ["invite", "last", "digest"],
            "notice_key": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_day": ["2", "4", "12", "13", "14"],
            "own_schedule": "on",
        },
    )
    assert saved.status_code == 200
    settings = app.state.store.partner(partner["id"])["notice_settings"]
    by_key = {item["key"]: item for item in settings["notices"]}
    assert settings["weekend_shift"] is True
    assert by_key["invite"]["day"] == 2
    assert by_key["reminder"]["enabled"] is False
    assert by_key["last"]["day"] == 12
    assert by_key["overdue"]["enabled"] is False
    assert settings["schedule_override"] is True


def test_global_notice_dates_apply_until_a_partner_overrides_them(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store

    page = client.get("/admin/settings")
    assert "Notification dates" in page.text
    saved = client.post(
        "/admin/settings/schedule",
        data={
            "notice_on": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_key": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_day": ["8", "9", "11", "12", "13"],
        },
    )
    assert saved.status_code == 200
    assert "Settings saved." in saved.text

    demo = store.partner_by_token("demo-partner")
    inherited = {item["key"]: item["day"] for item in demo["notice_settings"]["notices"]}
    assert inherited["invite"] == 8
    assert demo["notice_settings"]["schedule_override"] is False

    client.post(
        "/admin/partners",
        data={"name": "Exception App", "legal_name": "Exception LLC", "rate": "10"},
    )
    partner = next(row for row in store.partners() if row["name"] == "Exception App")
    client.post(
        f"/admin/partners/{partner['id']}/edit",
        data={
            "name": "Exception App",
            "legal_name": "Exception LLC",
            "legal_country": "Estonia",
            "altegio_legal_entity": "Altegio KFT",
            "contact_name": "Eva",
            "email": "eva@example.com",
            "telegram": "",
            "report_channel": "email",
            "application_url": "https://app.alteg.io/app/77",
            "application_comment": "",
            "rate": "10",
            "currency": "EUR",
            "notice_on": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_key": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_day": ["2", "4", "5", "6", "7"],
            "own_schedule": "on",
        },
    )
    personal = store.partner(partner["id"])["notice_settings"]
    assert personal["schedule_override"] is True
    assert next(item["day"] for item in personal["notices"] if item["key"] == "invite") == 2

    client.post(
        "/admin/settings/schedule",
        data={
            "notice_on": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_key": ["invite", "reminder", "last", "digest", "overdue"],
            "notice_day": ["3", "4", "5", "6", "7"],
        },
    )
    demo = store.partner_by_token("demo-partner")
    assert next(item["day"] for item in demo["notice_settings"]["notices"] if item["key"] == "invite") == 3
    personal = store.partner(partner["id"])["notice_settings"]
    assert next(item["day"] for item in personal["notices"] if item["key"] == "invite") == 2


def test_notice_texts_are_shared_across_partners(tmp_path: Path, monkeypatch):
    monkeypatch.setattr("revshare.web.telegram_token", lambda: "")
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store

    page = client.get("/admin/settings?language=ru")
    assert page.status_code == 200
    assert "Settings" in page.text
    assert "Telegram delivery" in page.text
    assert "Telegram bot is not connected" in page.text
    assert "Notification texts" in page.text
    assert "Rates to the dollar" in page.text
    assert "Здравствуйте" in page.text
    home = client.get("/")
    assert "Telegram delivery" not in home.text
    redirected = client.get("/admin/notices?language=ru", follow_redirects=False)
    assert redirected.status_code == 303
    assert "/admin/settings" in redirected.headers["location"]

    saved = client.post(
        "/admin/notices",
        data={
            "language": "ru",
            "notice_key": ["invite", "reminder", "last", "overdue"],
            "notice_message": [
                "Привет! Отчёт по {name} за {period} нужен до {deadline}. {link}",
                "",
                "",
                "",
            ],
        },
    )
    assert saved.status_code == 200
    assert "Settings saved." in saved.text

    partner = store.partner_by_token("demo-partner")
    store.update_partner(
        partner_id=partner["id"],
        name=partner["name"],
        legal_name="",
        legal_country="",
        altegio_legal_entity="Open HC",
        contact_name="",
        email="demo@example.com",
        telegram="",
        report_channel="email",
        rate=0.3,
        currency="USD",
        notice_settings={"language": "ru"},
    )
    refreshed = store.partner(partner["id"])["notice_settings"]
    invite = next(item for item in refreshed["notices"] if item["key"] == "invite")
    assert invite["message"].startswith("Привет! Отчёт по")
    reminder = next(item for item in refreshed["notices"] if item["key"] == "reminder")
    assert reminder["message"].startswith("Напоминаем")


def test_saving_snapshot_rechecks_already_submitted_report(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    registry = tmp_path / "registry.csv"
    registry.write_text("Location ID,Turnover\n101,$400\n", encoding="utf-8")
    with registry.open("rb") as handle:
        preview = client.post(
            "/submit/demo-partner/parse",
            data={"period": "2026-08"},
            files={"attachment": ("registry.csv", handle, "text/csv")},
        )
    draft_token = str(preview.url).rstrip("/").split("/")[-1]
    client.post(
        f"/submit/demo-partner/confirm/{draft_token}",
        data={
            "turnover": "400",
            "currency": "USD",
            "reported_share": "120",
            "location_ids": "101",
            "comment": "",
        },
    )
    report_id = store.report_id_for(partner["id"], "2026-08")
    assert store.report(report_id)["claims"] == []

    saved = client.post(
        "/admin/usage",
        data={
            "partner_id": str(partner["id"]),
            "period": "2026-08",
            "connected_ids": "101\n102",
            "active_ids": "101\n102",
            "payment_ids": "102",
        },
    )
    assert saved.status_code == 200
    assert saved.url.path == f"/reports/{report_id}"
    assert "No partner explanation yet" in saved.text
    assert (
        "A payment callback exists, and this Location ID is still missing from the report."
        in saved.text
    )
    report = store.report(report_id)
    assert report["status"] == "requires_review"
    assert report["missing_active"] == ["102"]


def test_json_snapshot_import_and_accrual_queue(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    partner = store.partner_by_token("demo-partner")
    payload = {
        "period": "2026-08",
        "partners": [
            {
                "token": "demo-partner",
                "connected": ["101", "102"],
                "active": ["101"],
                "payment": ["101"],
            }
        ],
    }
    imported = client.post(
        "/admin/settings/altegio/import",
        data={"period": "2026-08"},
        files={"attachment": ("snap.json", json.dumps(payload).encode(), "application/json")},
        follow_redirects=False,
    )
    assert imported.status_code == 303
    assert imported.headers["location"].endswith("imported=1#altegio")
    connected, active = store.usage(partner["id"], "2026-08")
    assert connected == {"101", "102"}
    assert active == {"101"}
    evidence = store.usage_import(partner["id"], "2026-08")
    assert evidence["source"] == "json"

    report_id, _ = store.submit_report(
        partner_id=partner["id"],
        period="2026-08",
        turnover=100,
        currency="USD",
        reported_share=30,
        location_ids=["101"],
        comment="",
        attachment=None,
    )
    store.approve_report(report_id)
    ready_page = client.get("/admin/accrual?period=2026-08")
    assert "Send selected to Google Sheet" in ready_page.text
    queued = client.post(
        "/admin/accrual/queue",
        data={"period": "2026-08", "report_ids": str(report_id)},
        follow_redirects=False,
    )
    assert queued.status_code == 303
    assert "queued=1" in queued.headers["location"]
    report = store.report(report_id)
    assert report["workflow_status"] == "posted"
    page = client.get("/admin/accrual?period=2026-08")
    assert page.status_code == 200
    assert "Partnerships Altegio" in page.text
    assert "Demo Integration Hub App" in page.text
    assert "Write queue" in page.text
    assert "100.00 USD" in page.text
    assert "30.00 USD" in page.text
    history = client.get(f"/reports/{report_id}")
    assert "Queued for Partnerships Altegio" in history.text
    notes = [event["note"] for event in store.report(report_id)["status_log"]]
    assert "queued" in notes


def test_accrual_bulk_queue_with_checkboxes(tmp_path: Path):
    app = create_app(tmp_path / "test.db")
    client = TestClient(app)
    store: Store = app.state.store
    first = store.partner_by_token("demo-partner")
    second_id = store.add_partner(
        name="Alpha App",
        legal_name="Alpha LLC",
        legal_country="Kazakhstan",
        altegio_legal_entity="",
        contact_name="",
        email="",
        telegram="",
        report_channel="email",
        rate=0.15,
        currency="USD",
        token="alpha-token",
    )
    report_ids = []
    for partner_id, location_id, turnover, share in (
        (first["id"], "101", 100, 30),
        (second_id, "102", 200, 60),
    ):
        store.replace_usage(
            partner_id, "2026-08", connected_ids=[location_id], active_ids=[location_id]
        )
        report_id, _ = store.submit_report(
            partner_id=partner_id,
            period="2026-08",
            turnover=turnover,
            currency="USD",
            reported_share=share,
            location_ids=[location_id],
            comment="",
            attachment=None,
        )
        store.approve_report(report_id)
        report_ids.append(report_id)

    ready = store.accrual_queue("2026-08")
    assert len(ready) == 2

    page = client.get("/admin/accrual?period=2026-08")
    assert page.status_code == 200
    assert 'name="report_ids"' in page.text
    assert "Send selected to Google Sheet" in page.text
    assert 'data-accrual-select-all' in page.text

    queued = client.post(
        "/admin/accrual/queue",
        data={"period": "2026-08", "report_ids": report_ids},
        follow_redirects=False,
    )
    assert queued.status_code == 303
    assert "queued=2" in queued.headers["location"]
    assert not store.accrual_queue("2026-08")
    assert len(store.sheet_writes("2026-08")) == 2
